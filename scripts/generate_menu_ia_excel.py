# -*- coding: utf-8 -*-
import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel():
    json_path = os.path.join(os.path.dirname(__file__), "..", "data", "menu_ia_data.json")
    with open(json_path, "r", encoding="utf-8") as f:
        meta = json.load(f)

    wb = openpyxl.Workbook()

    # Font definitions
    font_name = "\ub9d1\uc740 \uace0\ub515" # Malgun Gothic
    font_title = Font(name=font_name, size=15, bold=True, color="1B365D")
    font_subtitle = Font(name=font_name, size=9.5, italic=True, color="555555")
    font_header = Font(name=font_name, size=10.5, bold=True, color="FFFFFF")
    font_body = Font(name=font_name, size=9.5, color="222222")
    font_body_bold = Font(name=font_name, size=9.5, bold=True, color="222222")
    font_code = Font(name="Consolas", size=9, color="1B365D", bold=True)

    # Fill definitions
    fill_header_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    fill_header_blue = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
    fill_header_green = PatternFill(start_color="1E7E34", end_color="1E7E34", fill_type="solid")
    fill_header_gray = PatternFill(start_color="3B4252", end_color="3B4252", fill_type="solid")

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Border definitions
    thin_side = Side(style="thin", color="D1D5DB")
    medium_top = Side(style="medium", color="1B365D")
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_header = Border(left=thin_side, right=thin_side, top=medium_top, bottom=medium_top)

    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # -------------------------------------------------------------
    # 1. Sheet 1: 메뉴구조도_총괄(IA)
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = meta['sheet1_name']
    ws1.views.sheetView[0].showGridLines = True
    ws1.sheet_properties.tabColor = "1B365D"

    ws1.cell(row=2, column=2, value=meta['doc_title']).font = font_title
    ws1.cell(row=3, column=2, value=meta['doc_subtitle']).font = font_subtitle
    ws1.row_dimensions[2].height = 28
    ws1.row_dimensions[3].height = 18

    start_r = 5
    for c_idx, h in enumerate(meta['sheet1_headers'], start=2):
        cell = ws1.cell(row=start_r, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = align_header
        cell.border = border_header
    ws1.row_dimensions[start_r].height = 28

    curr_r = start_r + 1
    for idx, r_data in enumerate(meta['sheet1_rows'], start=1):
        row_fill = fill_zebra if idx % 2 == 0 else fill_white
        ws1.cell(row=curr_r, column=2, value=idx).alignment = align_center
        ws1.cell(row=curr_r, column=3, value=r_data[0]).alignment = align_center # ID
        ws1.cell(row=curr_r, column=4, value=r_data[1]).alignment = align_center # 대분류
        ws1.cell(row=curr_r, column=5, value=r_data[2]).alignment = align_center # 중분류
        ws1.cell(row=curr_r, column=6, value=r_data[3]).alignment = align_left   # 소분류
        ws1.cell(row=curr_r, column=7, value=r_data[4]).alignment = align_left   # 주요기능
        ws1.cell(row=curr_r, column=8, value=r_data[5]).alignment = align_left   # 컴포넌트
        ws1.cell(row=curr_r, column=9, value=r_data[6]).alignment = align_center # 권한
        ws1.cell(row=curr_r, column=10, value=r_data[7]).alignment = align_left  # 연계데이터
        ws1.cell(row=curr_r, column=11, value=r_data[8]).alignment = align_left  # AI기술
        ws1.cell(row=curr_r, column=12, value=r_data[9]).alignment = align_left  # 비고

        for c_idx in range(2, 13):
            cell = ws1.cell(row=curr_r, column=c_idx)
            cell.font = font_body
            cell.fill = row_fill
            cell.border = border_cell
            if c_idx == 3:
                cell.font = font_code
            elif c_idx == 4:
                cell.font = font_body_bold

        ws1.row_dimensions[curr_r].height = 25
        curr_r += 1

    # -------------------------------------------------------------
    # 2. Sheet 2: 워크스페이스별_상세명세
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title=meta['sheet2_name'])
    ws2.views.sheetView[0].showGridLines = True
    ws2.sheet_properties.tabColor = "2E5B88"

    ws2.cell(row=2, column=2, value="\u0036\ub300 \ud575\uc2ec \uc6cc\ud06c\uc2a4\ud398\uc774\uc2a4\ubcc4 \uae30\ub2a5 \uba85\uc138 \ubc0f \uc785\ucd9c\ub825 \uc0b0\ucd9c\ubb3c").font = font_title
    ws2.cell(row=3, column=2, value="\uc0ac\uc6a9\uc790 \uc791\uc5c5 \ud750\ub984(Workflow) \ubc0f \uc0b0\ucd9c\ubb3c \ud30c\uc77c \uc720\ud615 \uc815\uc758").font = font_subtitle
    ws2.row_dimensions[2].height = 28
    ws2.row_dimensions[3].height = 18

    for c_idx, h in enumerate(meta['sheet2_headers'], start=2):
        cell = ws2.cell(row=start_r, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_blue
        cell.alignment = align_header
        cell.border = border_header
    ws2.row_dimensions[start_r].height = 28

    curr_r = start_r + 1
    for idx, r_data in enumerate(meta['sheet2_rows'], start=1):
        row_fill = fill_zebra if idx % 2 == 0 else fill_white
        ws2.cell(row=curr_r, column=2, value=idx).alignment = align_center
        ws2.cell(row=curr_r, column=3, value=r_data[0]).alignment = align_center
        ws2.cell(row=curr_r, column=4, value=r_data[1]).alignment = align_left
        ws2.cell(row=curr_r, column=5, value=r_data[2]).alignment = align_left
        ws2.cell(row=curr_r, column=6, value=r_data[3]).alignment = align_left
        ws2.cell(row=curr_r, column=7, value=r_data[4]).alignment = align_left
        ws2.cell(row=curr_r, column=8, value=r_data[5]).alignment = align_left
        ws2.cell(row=curr_r, column=9, value=r_data[6]).alignment = align_center

        for c_idx in range(2, 10):
            cell = ws2.cell(row=curr_r, column=c_idx)
            cell.font = font_body
            cell.fill = row_fill
            cell.border = border_cell
            if c_idx == 3:
                cell.font = font_body_bold

        ws2.row_dimensions[curr_r].height = 32
        curr_r += 1

    # -------------------------------------------------------------
    # 3. Sheet 3: 연계데이터_및_AI기능
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title=meta['sheet3_name'])
    ws3.views.sheetView[0].showGridLines = True
    ws3.sheet_properties.tabColor = "1E7E34"

    ws3.cell(row=2, column=2, value="\uacf5\uacf5\ubcf4\uac74\uc758\ub8cc \uc5f0\uacc4 \uacf5\uacf5\ub370\uc774\ud130\uc14b \ubc0f AI \ubaa8\ub378 \uba85\uc138").font = font_title
    ws3.cell(row=3, column=2, value="\uad6d\uac00 \uacf5\uacf5\ubcf4\uac74\uc758\ub8cc \uc6d0\ucc9c \ub370\uc774\ud130\uc14b, OpenAPI \uaddc\uaca9 \ubc0f AI \uc5d4\uc9c4 \uc0ac\uc591").font = font_subtitle
    ws3.row_dimensions[2].height = 28
    ws3.row_dimensions[3].height = 18

    for c_idx, h in enumerate(meta['sheet3_headers'], start=2):
        cell = ws3.cell(row=start_r, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_green
        cell.alignment = align_header
        cell.border = border_header
    ws3.row_dimensions[start_r].height = 28

    curr_r = start_r + 1
    for idx, r_data in enumerate(meta['sheet3_rows'], start=1):
        row_fill = fill_zebra if idx % 2 == 0 else fill_white
        ws3.cell(row=curr_r, column=2, value=idx).alignment = align_center
        ws3.cell(row=curr_r, column=3, value=r_data[0]).alignment = align_left
        ws3.cell(row=curr_r, column=4, value=r_data[1]).alignment = align_center
        ws3.cell(row=curr_r, column=5, value=r_data[2]).alignment = align_left
        ws3.cell(row=curr_r, column=6, value=r_data[3]).alignment = align_left
        ws3.cell(row=curr_r, column=7, value=r_data[4]).alignment = align_center
        ws3.cell(row=curr_r, column=8, value=r_data[5]).alignment = align_left
        ws3.cell(row=curr_r, column=9, value=r_data[6]).alignment = align_left

        for c_idx in range(2, 10):
            cell = ws3.cell(row=curr_r, column=c_idx)
            cell.font = font_body
            cell.fill = row_fill
            cell.border = border_cell
            if c_idx == 3:
                cell.font = font_body_bold

        ws3.row_dimensions[curr_r].height = 26
        curr_r += 1

    # -------------------------------------------------------------
    # 4. Sheet 4: 문서개정이력
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title=meta['sheet4_name'])
    ws4.views.sheetView[0].showGridLines = True
    ws4.sheet_properties.tabColor = "3B4252"

    ws4.cell(row=2, column=2, value="\uacf5\uacf5\ubcf4\uac74\uc758\ub8cc \ube45\ub370\uc774\ud130 \ud50c\ub7ab\ud3fc \uba54\ub274\uad6c\uc870\ub3c4 \uac1c\uc815 \uc774\ub825").font = font_title
    ws4.cell(row=3, column=2, value="\ud50c\ub7ab\ud3fc \uae30\ub2a5 \ubcc0\uacbd \ubc0f \ubc84\uc804 \uc5c5\ub370\uc774\ud2b8 \uae30\ub85d \uad00\ub9ac").font = font_subtitle
    ws4.row_dimensions[2].height = 28
    ws4.row_dimensions[3].height = 18

    for c_idx, h in enumerate(meta['sheet4_headers'], start=2):
        cell = ws4.cell(row=start_r, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_gray
        cell.alignment = align_header
        cell.border = border_header
    ws4.row_dimensions[start_r].height = 28

    curr_r = start_r + 1
    for idx, r_data in enumerate(meta['sheet4_rows'], start=1):
        row_fill = fill_zebra if idx % 2 == 0 else fill_white
        ws4.cell(row=curr_r, column=2, value=r_data[0]).alignment = align_center # 버젼
        ws4.cell(row=curr_r, column=3, value=r_data[1]).alignment = align_center # 일자
        ws4.cell(row=curr_r, column=4, value=r_data[2]).alignment = align_center # 작성자
        ws4.cell(row=curr_r, column=5, value=r_data[3]).alignment = align_left   # 내용
        ws4.cell(row=curr_r, column=6, value=r_data[4]).alignment = align_center # 상태
        ws4.cell(row=curr_r, column=7, value=r_data[5]).alignment = align_left   # 비고

        for c_idx in range(2, 8):
            cell = ws4.cell(row=curr_r, column=c_idx)
            cell.font = font_body
            cell.fill = row_fill
            cell.border = border_cell
            if c_idx == 2:
                cell.font = font_code

        ws4.row_dimensions[curr_r].height = 26
        curr_r += 1

    # -------------------------------------------------------------
    # Formatting adjustments (Column widths & freeze panes)
    # -------------------------------------------------------------
    for sheet in wb.worksheets:
        sheet.column_dimensions['A'].width = 3
        sheet.freeze_panes = "B6"

        for col in sheet.columns:
            col_letter = get_column_letter(col[0].column)
            if col_letter == 'A':
                continue
            max_len = 0
            for cell in col:
                if cell.row < 5:
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    # Count wide chars (hangul, cjk) as 2
                    length = sum(2 if ord(char) > 127 else 1 for char in val_str)
                    if length > max_len:
                        max_len = length
            calc_w = max(max_len + 3, 11)
            if calc_w > 50:
                calc_w = 50
            sheet.column_dimensions[col_letter].width = calc_w

    # Fine-tune ws1 column widths
    ws1.column_dimensions['B'].width = 7   # 순번
    ws1.column_dimensions['C'].width = 13  # 메뉴ID
    ws1.column_dimensions['D'].width = 14  # 대분류
    ws1.column_dimensions['E'].width = 16  # 중분류
    ws1.column_dimensions['F'].width = 28  # 소분류
    ws1.column_dimensions['G'].width = 46  # 주요기능
    ws1.column_dimensions['H'].width = 25  # 컴포넌트
    ws1.column_dimensions['I'].width = 18  # 권한
    ws1.column_dimensions['J'].width = 28  # 연계데이터
    ws1.column_dimensions['K'].width = 26  # AI기술
    ws1.column_dimensions['L'].width = 24  # 비고

    # Save to user Downloads folder
    downloads_dir = os.path.expanduser("~/Downloads")
    # 파일명을 unicode-escape로 안전하게 구성: 공공보건의료_플랫폼_메뉴구조도_IA.xlsx
    file_name = "\uacf5\uacf5\ubcf4\uac74\uc758\ub8cc_\ud50c\ub7ab\ud3fc_\uba54\ub274\uad6c\uc870\ub3c4_IA.xlsx"
    target_path = os.path.join(downloads_dir, file_name)
    target_path_v2 = os.path.join(downloads_dir, '공공보건의료_플랫폼_메뉴구조도_IA_최신.xlsx')
    try:
        wb.save(target_path)
        print('Saved to original:', target_path)
    except PermissionError:
        wb.save(target_path_v2)
        print('Saved to v2 (since original is locked in Excel):', target_path_v2)
    print("SUCCESS: Target path code points:", [ord(c) for c in file_name])
    print("Saved file size:", os.path.getsize(target_path), "bytes")

if __name__ == "__main__":
    build_excel()
