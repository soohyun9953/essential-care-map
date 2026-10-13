# -*- coding: utf-8 -*-
import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_user_menu_excel():
    json_path = os.path.join(os.path.dirname(__file__), "..", "data", "user_menu_matrix_data.json")
    with open(json_path, "r", encoding="utf-8") as f:
        meta = json.load(f)

    wb = openpyxl.Workbook()

    font_name = "\ub9d1\uc740 \uace0\ub515"  # 맑은 고딕
    font_title = Font(name=font_name, size=15, bold=True, color="1B365D")
    font_subtitle = Font(name=font_name, size=9.5, italic=True, color="555555")
    font_header = Font(name=font_name, size=10.5, bold=True, color="FFFFFF")
    font_body = Font(name=font_name, size=9.5, color="222222")
    font_body_bold = Font(name=font_name, size=9.5, bold=True, color="1B365D")
    font_code = Font(name="Consolas", size=9, color="1B365D", bold=True)

    # User Group Fonts & Fills
    group_styles = {
        "\u2460": (Font(name=font_name, size=9.5, bold=True, color="0F5132"), PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")), # ① 대국민 (Green)
        "\u2461": (Font(name=font_name, size=9.5, bold=True, color="084298"), PatternFill(start_color="CFE2FF", end_color="CFE2FF", fill_type="solid")), # ② 의료기관 (Blue)
        "\u2462": (Font(name=font_name, size=9.5, bold=True, color="3D0A91"), PatternFill(start_color="E0CFFC", end_color="E0CFFC", fill_type="solid")), # ③ 지역공공 (Purple)
        "\u2463": (Font(name=font_name, size=9.5, bold=True, color="664D03"), PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")), # ④ 시스템담당 (Amber)
    }

    # Matrix Cell Styles
    font_star = Font(name=font_name, size=9.5, bold=True, color="1E40AF")
    fill_star = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    font_circle = Font(name=font_name, size=9, color="334155")
    fill_circle = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_dash = Font(name=font_name, size=9, color="94A3B8")

    # Header Fills
    fill_header_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    fill_header_blue = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
    fill_header_teal = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    fill_header_purple = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_side = Side(style="thin", color="CBD5E1")
    medium_top = Side(style="medium", color="1B365D")
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_header = Border(left=thin_side, right=thin_side, top=medium_top, bottom=medium_top)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # =========================================================================
    # Sheet 1: 1_사용자별_메뉴구성_총괄
    # =========================================================================
    ws1 = wb.active
    ws1.title = meta["s1_name"]
    ws1.views.sheetView[0].showGridLines = True
    ws1.sheet_properties.tabColor = "1B365D"

    ws1.cell(row=2, column=2, value=meta["doc_title"]).font = font_title
    ws1.cell(row=3, column=2, value=meta["doc_subtitle"]).font = font_subtitle
    ws1.row_dimensions[2].height = 28
    ws1.row_dimensions[3].height = 18

    start_r = 5
    for c_idx, h in enumerate(meta["s1_headers"], start=2):
        cell = ws1.cell(row=start_r, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = align_header
        cell.border = border_header
    ws1.row_dimensions[start_r].height = 30

    curr_r = start_r + 1
    for idx, r_data in enumerate(meta["s1_rows"], start=1):
        row_fill = fill_zebra if idx % 2 == 0 else fill_white
        ws1.cell(row=curr_r, column=2, value=idx).alignment = align_center
        ws1.cell(row=curr_r, column=3, value=r_data[0]).alignment = align_center  # 사용자 구분
        ws1.cell(row=curr_r, column=4, value=r_data[1]).alignment = align_left    # 세부 소속
        ws1.cell(row=curr_r, column=5, value=r_data[2]).alignment = align_left    # 핵심 업무 목적
        ws1.cell(row=curr_r, column=6, value=r_data[3]).alignment = align_center  # 상위 워크스페이스
        ws1.cell(row=curr_r, column=7, value=r_data[4]).alignment = align_left    # 세부 메뉴명
        ws1.cell(row=curr_r, column=8, value=r_data[5]).alignment = align_left    # 주요 제공 기능
        ws1.cell(row=curr_r, column=9, value=r_data[6]).alignment = align_left    # 핵심 산출물
        ws1.cell(row=curr_r, column=10, value=r_data[7]).alignment = align_center # 접근 수준

        for c_idx in range(2, 11):
            cell = ws1.cell(row=curr_r, column=c_idx)
            cell.font = font_body
            cell.fill = row_fill
            cell.border = border_cell
            if c_idx == 3:
                prefix = str(r_data[0])[:1]
                if prefix in group_styles:
                    cell.font, cell.fill = group_styles[prefix]
            elif c_idx in (6, 7):
                cell.font = font_body_bold

        ws1.row_dimensions[curr_r].height = 32
        curr_r += 1

    # =========================================================================
    # Sheet 2: 2_메뉴별_사용자권한_매트릭스
    # =========================================================================
    ws2 = wb.create_sheet(title=meta["s2_name"])
    ws2.views.sheetView[0].showGridLines = True
    ws2.sheet_properties.tabColor = "2E5B88"

    ws2.cell(row=2, column=2, value="\ud50c\ub7ab\ud3fc \uc804\uccb4 \uba54\ub274(32\uac1c) \ubc0f 4\ub300 \uc0ac\uc6a9\uc790 \uadf8\ub8f9\ubcc4 \ud65c\uc6a9 \ub9e4\ud2b8\ub9ad\uc2a4").font = font_title
    ws2.cell(row=3, column=2, value="\ubc94\ub840: \u2605 \uc8fc\uc0ac\uc6a9(\ud575\uc2ec \uc5c5\ubb34 \uc218\ud589)  |  \u25cb \uc5f4\ub78c\u00b7\ucc38\uc870(\ubcf4\uc870 \ud65c\uc6a9)  |  - \ud574\ub2f9\uc5c6\uc74c(\ube44\ub178\ucd9c \ub610\ub294 \uad8c\ud55c\uc81c\ud55c)").font = font_subtitle
    ws2.row_dimensions[2].height = 28
    ws2.row_dimensions[3].height = 18

    for c_idx, h in enumerate(meta["s2_headers"], start=2):
        cell = ws2.cell(row=start_r, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_blue
        cell.alignment = align_header
        cell.border = border_header
    ws2.row_dimensions[start_r].height = 34

    curr_r = start_r + 1
    for idx, r_data in enumerate(meta["s2_rows"], start=1):
        row_fill = fill_zebra if idx % 2 == 0 else fill_white
        ws2.cell(row=curr_r, column=2, value=idx).alignment = align_center
        ws2.cell(row=curr_r, column=3, value=r_data[0]).alignment = align_center  # 메뉴코드
        ws2.cell(row=curr_r, column=4, value=r_data[1]).alignment = align_center  # 상위 워크스페이스
        ws2.cell(row=curr_r, column=5, value=r_data[2]).alignment = align_left    # 세부 메뉴
        ws2.cell(row=curr_r, column=6, value=r_data[3]).alignment = align_center  # ① 대국민
        ws2.cell(row=curr_r, column=7, value=r_data[4]).alignment = align_center  # ② 의료기관
        ws2.cell(row=curr_r, column=8, value=r_data[5]).alignment = align_center  # ③ 지역공공
        ws2.cell(row=curr_r, column=9, value=r_data[6]).alignment = align_center  # ④ 시스템담당
        ws2.cell(row=curr_r, column=10, value=r_data[7]).alignment = align_center # 연계 시스템
        ws2.cell(row=curr_r, column=11, value=r_data[8]).alignment = align_left   # 비고

        for c_idx in range(2, 12):
            cell = ws2.cell(row=curr_r, column=c_idx)
            cell.font = font_body
            cell.fill = row_fill
            cell.border = border_cell
            if c_idx == 3:
                cell.font = font_code
            elif c_idx == 4:
                cell.font = font_body_bold
            elif c_idx in (6, 7, 8, 9):
                val = str(cell.value or "")
                if "\u2605" in val:
                    cell.font = font_star
                    cell.fill = fill_star
                elif "\u25cb" in val:
                    cell.font = font_circle
                    cell.fill = fill_circle
                else:
                    cell.font = font_dash

        ws2.row_dimensions[curr_r].height = 24
        curr_r += 1

    # =========================================================================
    # Sheet 3: 3_사용자별_업무시나리오_동선
    # =========================================================================
    ws3 = wb.create_sheet(title=meta["s3_name"])
    ws3.views.sheetView[0].showGridLines = True
    ws3.sheet_properties.tabColor = "0F766E"

    ws3.cell(row=2, column=2, value="4\ub300 \uc0ac\uc6a9\uc790 \uc720\ud615\ubcc4 \ud45c\uc900 \uc5c5\ubb34 \uc2dc\ub098\ub9ac\uc624 \ubc0f \uba54\ub274 \uc774\ub3d9 \ub3d9\uc120 (User Journey)").font = font_title
    ws3.cell(row=3, column=2, value="\uc811\uc18d \uc9c4\uc785\uc810\ubd80\ud130 \ud0d0\uc0c9\u00b7\uc9c4\ub2e8 \u2794 \ube44\uad50\u00b7\ubd84\uc11d \u2794 AI \uc2dc\ubba4\ub808\uc774\uc158 \u2794 \ucd5c\uc885 \uc0b0\ucd9c\ubb3c \ub3c4\ucd9c\uae4c\uc9c0\uc758 End-to-End \ud750\ub984").font = font_subtitle
    ws3.row_dimensions[2].height = 28
    ws3.row_dimensions[3].height = 18

    for c_idx, h in enumerate(meta["s3_headers"], start=2):
        cell = ws3.cell(row=start_r, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_teal
        cell.alignment = align_header
        cell.border = border_header
    ws3.row_dimensions[start_r].height = 30

    curr_r = start_r + 1
    for idx, r_data in enumerate(meta["s3_rows"], start=1):
        row_fill = fill_zebra if idx % 2 == 0 else fill_white
        for offset, val in enumerate(r_data):
            c_idx = 2 + offset
            cell = ws3.cell(row=curr_r, column=c_idx, value=val)
            cell.font = font_body
            cell.fill = row_fill
            cell.border = border_cell
            cell.alignment = align_center if offset in (0, 1, 2) else align_left
            if offset == 0:
                prefix = str(val)[:1]
                if prefix in group_styles:
                    cell.font, cell.fill = group_styles[prefix]
            elif offset == 7:
                cell.font = font_body_bold

        ws3.row_dimensions[curr_r].height = 48
        curr_r += 1

    # =========================================================================
    # Sheet 4: 4_8대시스템_사용자매핑
    # =========================================================================
    ws4 = wb.create_sheet(title=meta["s4_name"])
    ws4.views.sheetView[0].showGridLines = True
    ws4.sheet_properties.tabColor = "4C1D95"

    ws4.cell(row=2, column=2, value="\uacf5\uacf5\uc758\ub8cc 8\ub300 \ud575\uc2ec \uc2dc\uc2a4\ud15c \ubc0f 4\ub300 \uc0ac\uc6a9\uc790 \uadf8\ub8f9 \ub9e4\ud551\ud45c").font = font_title
    ws4.cell(row=3, column=2, value="\ud50c\ub7ab\ud3fc \ub0b4 [\ub2f4\ub2f9\uc790\ubcc4 \uac80\ud1a0 \uac00\uc774\ub4dc]\uc5d0 \uc815\uc758\ub41c \uacf5\uacf5\uc758\ub8cc \uc2dc\uc2a4\ud15c\ubcc4 \uba54\ub274 \uacbd\ub85c \ubc0f \ud0c0\uac9f \uc0ac\uc6a9\uc790").font = font_subtitle
    ws4.row_dimensions[2].height = 28
    ws4.row_dimensions[3].height = 18

    for c_idx, h in enumerate(meta["s4_headers"], start=2):
        cell = ws4.cell(row=start_r, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_purple
        cell.alignment = align_header
        cell.border = border_header
    ws4.row_dimensions[start_r].height = 30

    curr_r = start_r + 1
    for idx, r_data in enumerate(meta["s4_rows"], start=1):
        row_fill = fill_zebra if idx % 2 == 0 else fill_white
        for offset, val in enumerate(r_data):
            c_idx = 2 + offset
            cell = ws4.cell(row=curr_r, column=c_idx, value=val)
            cell.font = font_body
            cell.fill = row_fill
            cell.border = border_cell
            cell.alignment = align_center if offset in (0, 3, 6) else align_left
            if offset in (1, 3):
                cell.font = font_body_bold

        ws4.row_dimensions[curr_r].height = 32
        curr_r += 1

    # Column widths & Freeze Panes
    for sheet in wb.worksheets:
        sheet.column_dimensions["A"].width = 3
        sheet.freeze_panes = "B6"

    # Sheet 1 widths
    s1_widths = {"B": 6, "C": 22, "D": 28, "E": 34, "F": 20, "G": 34, "H": 48, "I": 30, "J": 20}
    for col_l, w in s1_widths.items():
        ws1.column_dimensions[col_l].width = w

    # Sheet 2 widths
    s2_widths = {"B": 6, "C": 11, "D": 18, "E": 38, "F": 15, "G": 16, "H": 17, "I": 17, "J": 22, "K": 36}
    for col_l, w in s2_widths.items():
        ws2.column_dimensions[col_l].width = w

    # Sheet 3 widths
    s3_widths = {"B": 22, "C": 26, "D": 24, "E": 34, "F": 34, "G": 36, "H": 34, "I": 32}
    for col_l, w in s3_widths.items():
        ws3.column_dimensions[col_l].width = w

    # Sheet 4 widths
    s4_widths = {"B": 12, "C": 26, "D": 36, "E": 30, "F": 30, "G": 46, "H": 26}
    for col_l, w in s4_widths.items():
        ws4.column_dimensions[col_l].width = w

    downloads_dir = os.path.expanduser("~/Downloads")
    # 공공보건의료_플랫폼_사용자별_메뉴구성도.xlsx
    file_name = "\uacf5\uacf5\ubcf4\uac74\uc758\ub8cc_\ud50c\ub7ab\ud3fc_\uc0ac\uc6a9\uc790\ubcc4_\uba54\ub274\uad6c\uc131\ub3c4.xlsx"
    target_path = os.path.join(downloads_dir, file_name)
    target_path_v2 = os.path.join(downloads_dir, "\uacf5\uacf5\ubcf4\uac74\uc758\ub8cc_\ud50c\ub7ab\ud3fc_\uc0ac\uc6a9\uc790\ubcc4_\uba54\ub274\uad6c\uc131\ub3c4_\ucd5c\uc2e0.xlsx")

    try:
        wb.save(target_path)
        print("SUCCESS: Saved to:", target_path)
        print("File size:", os.path.getsize(target_path), "bytes")
    except PermissionError:
        wb.save(target_path_v2)
        print("SUCCESS: Saved to fallback:", target_path_v2)
        print("File size:", os.path.getsize(target_path_v2), "bytes")

if __name__ == "__main__":
    build_user_menu_excel()
